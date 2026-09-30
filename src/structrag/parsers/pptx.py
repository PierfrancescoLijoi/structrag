"""PPTX parser: one level-1 section per slide, title placeholder = heading, notes included."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from ..ir import LIST_ITEM, PARA, TABLE, Block, ParsedDoc

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}


def _shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes)
        else:
            yield shape


def _is_title(shape) -> bool:
    return shape.is_placeholder and shape.placeholder_format.type in TITLE_TYPES


def _slide_blocks(slide, number: int) -> list[Block]:
    loc = f"slide {number}"
    shapes = sorted(_shapes(slide.shapes), key=lambda s: (s.top or 0, s.left or 0))
    title = next((s.text_frame.text.strip() for s in shapes if _is_title(s) and s.has_text_frame), "")
    blocks = [Block(title or f"Slide {number}", level_hint=1, style="slide-title", loc=loc,
                    meta={"untitled": not title})]
    for shape in shapes:
        if _is_title(shape):
            continue
        if shape.has_text_frame:
            pars = shape.text_frame.paragraphs
            for par in pars:
                text = "".join(r.text for r in par.runs).strip()
                if text:
                    kind = LIST_ITEM if par.level or len(pars) > 1 else PARA
                    blocks.append(Block(text, kind, loc=loc))
        elif getattr(shape, "has_table", False) and shape.has_table:
            rows = tuple(tuple(" ".join(c.text.split()) for c in r.cells) for r in shape.table.rows)
            blocks.append(Block("\n".join(" | ".join(r) for r in rows), TABLE, rows=rows, loc=loc))
    if slide.has_notes_slide and (notes := slide.notes_slide.notes_text_frame.text.strip()):
        blocks.append(Block(f"Speaker notes: {notes}", PARA, loc=loc, meta={"notes": True}))
    return blocks


def parse(path: Path) -> ParsedDoc:
    prs = Presentation(str(path))
    blocks: list[Block] = []
    for number, slide in enumerate(prs.slides, 1):
        blocks.extend(_slide_blocks(slide, number))
    title = prs.core_properties.title or (blocks[0].text if blocks else path.stem)
    return ParsedDoc(title=title, format="pptx", blocks=tuple(blocks))
