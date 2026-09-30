"""XLSX parser: sheets -> regions (blank row/column gaps) -> tables with header detection.

Heuristics (cheap first, agent resolves the ambiguous ones later):
  * a sheet is level-1; every region gets a table block with its cell range
  * leading single-cell rows above a region are titles (level-2 heading)
  * header row = best-scoring of the first rows (text-only row followed by mixed types, bold bonus)
  * merged cells above the header are joined as "parent / child" column names
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..config import Settings
from ..ir import TABLE, Block, ParsedDoc

HEADER_SCAN_ROWS = 5
AMBIGUOUS_MARGIN = 0.15   # top-2 header scores closer than this => ask the agent
MIN_HEADER_SCORE = 0.5


def _fmt(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    if isinstance(value, float):
        return f"{value:.6g}"
    return str(value).strip()


def _is_number(text: str) -> bool:
    try:
        float(text.replace(",", ".").replace("%", "").replace("€", "").replace("$", ""))
        return True
    except ValueError:
        return False


def _spans(flags: list[bool]) -> list[tuple[int, int]]:
    """Maximal runs of True in flags as inclusive (start, end)."""
    spans, start = [], None
    for i, flag in enumerate(flags + [False]):
        if flag and start is None:
            start = i
        elif not flag and start is not None:
            spans.append((start, i - 1))
            start = None
    return spans


def _regions(grid: list[list[str]]) -> list[tuple[int, int, int, int]]:
    """Rectangles (r0, r1, c0, c1) separated by fully empty rows, then fully empty columns."""
    out = []
    for r0, r1 in _spans([any(row) for row in grid]):
        width = len(grid[0])
        col_used = [any(grid[r][c] for r in range(r0, r1 + 1)) for c in range(width)]
        out.extend((r0, r1, c0, c1) for c0, c1 in _spans(col_used))
    return out


def _header_scores(rows: list[list[str]], bold_rows: set[int]) -> list[float]:
    scores = []
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        filled = [c for c in row if c]
        if not filled or i + 1 >= len(rows):
            scores.append(0.0)
            continue
        text_share = sum(not _is_number(c) for c in filled) / len(filled)
        fill_share = len(filled) / len(row)
        below = [c for c in rows[i + 1] if c]
        below_numeric = (sum(_is_number(c) for c in below) / len(below)) if below else 0.0
        score = 0.45 * text_share + 0.25 * fill_share + 0.2 * below_numeric + (0.1 if i in bold_rows else 0.0)
        scores.append(score * (0.95 ** i))  # earlier rows win ties
    return scores


def _merged_header_prefix(ws, top: int, header_row: int, c0: int, c1: int, width: int) -> list[str]:
    """Names of merged cells sitting in the rows above the header, per column (1-based sheet coords)."""
    prefix = [""] * width
    for rng in ws.merged_cells.ranges:
        if rng.min_row >= top and rng.max_row < header_row and rng.min_col <= c1 + 1 and rng.max_col >= c0 + 1:
            label = _fmt(ws.cell(rng.min_row, rng.min_col).value)
            for col in range(max(rng.min_col, c0 + 1), min(rng.max_col, c1 + 1) + 1):
                prefix[col - c0 - 1] = label
    return prefix


def _region_blocks(ws, grid, region, bold_map: set[int], sheet: str) -> list[Block]:
    r0, r1, c0, c1 = region
    rows = [grid[r][c0:c1 + 1] for r in range(r0, r1 + 1)]
    blocks: list[Block] = []
    # Leading single-cell rows above a multi-row table are titles.
    while len(rows) > 2 and sum(bool(c) for c in rows[0]) == 1:
        blocks.append(Block(next(c for c in rows[0] if c), level_hint=2, style="xlsx-title", loc=sheet))
        rows, r0 = rows[1:], r0 + 1
    if not rows:
        return blocks
    bold_rows = {i for i in range(min(HEADER_SCAN_ROWS, len(rows))) if (r0 + i) in bold_map}
    scores = _header_scores(rows, bold_rows)
    best = max(range(len(scores)), key=scores.__getitem__) if scores else 0
    ranked = sorted(scores, reverse=True)
    ambiguous = len(rows) > 1 and (ranked[0] < MIN_HEADER_SCORE
                                   or (len(ranked) > 1 and ranked[0] - ranked[1] < AMBIGUOUS_MARGIN))
    width = c1 - c0 + 1
    header = rows[best]
    prefix = _merged_header_prefix(ws, r0 + 1, r0 + best + 1, c0, c1, width)
    header = [f"{p} / {h}" if p and h and p != h else h for p, h in zip(prefix, header)]
    header = [h or f"col{get_column_letter(c0 + i + 1)}" for i, h in enumerate(header)]
    all_rows = tuple(tuple(r) for r in rows)   # keep every row so `header_row` stays a valid index
    cell_range = f"{sheet}!{get_column_letter(c0 + 1)}{r0 + 1}:{get_column_letter(c1 + 1)}{r1 + 1}"
    text = "\n".join(" | ".join(r) for r in all_rows)
    blocks.append(Block(text, TABLE, rows=all_rows, loc=cell_range,
                        meta={"range": cell_range, "header_row": best, "column_names": header,
                              "header_ambiguous": ambiguous, "sheet": sheet}))
    return blocks


def parse(path: Path, settings: Settings | None = None) -> ParsedDoc:
    settings = settings or Settings()
    wb = load_workbook(str(path), data_only=True)
    blocks: list[Block] = []
    warnings: list[str] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" or not ws.max_row:
            continue
        max_rows = max(1, settings.max_xlsx_cells // max(1, ws.max_column))
        if ws.max_row > max_rows:
            warnings.append(f"sheet '{ws.title}' truncated to {max_rows} rows")
        grid = [[_fmt(c.value) for c in row] for row in ws.iter_rows(max_row=max_rows)]
        # Merged ranges other than the top-left are empty in the grid: that is what we want for data rows.
        bold_map = {r for r, row in enumerate(grid)
                    if any(ws.cell(r + 1, c + 1).font.bold for c, v in enumerate(row) if v)}
        blocks.append(Block(ws.title, level_hint=1, style="sheet-name", loc=ws.title))
        for region in _regions(grid):
            blocks.extend(_region_blocks(ws, grid, region, bold_map, ws.title))
    title = wb.properties.title or path.stem
    return ParsedDoc(title=title, format="xlsx", blocks=tuple(blocks), warnings=tuple(warnings))
