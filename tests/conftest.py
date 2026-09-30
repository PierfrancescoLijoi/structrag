"""Shared fixtures: small real documents in every supported format, generated on the fly."""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from structrag.app import Services, build_services
from structrag.config import Settings
from structrag.llm.fake import FakeLLM

LOREM = ("The quick brown fox jumps over the lazy dog while the committee reviews the quarterly "
         "results and the auditors verify every figure in the annual statement carefully. ") * 3


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(data_dir=tmp_path / "data", profiles_dir=tmp_path / "profiles", embedder="hash")


def heading_oracle(*titles: str):
    """Fake agent: answers A (heading) when the candidate line is one of `titles`, else B."""
    def choose(user: str, labels: list[str]) -> dict[str, float]:
        line = user.split("CANDIDATE LINE: ", 1)[1].split("\n", 1)[0]
        yes = any(line.startswith(t) for t in titles)
        return {"A": 0.95, "B": 0.05} if yes else {"A": 0.03, "B": 0.97}
    return choose


@pytest.fixture
def services(settings: Settings) -> Services:
    return build_services(settings, llm=FakeLLM(reply="Answer [1]"))


@pytest.fixture
def md_file(tmp_path: Path) -> Path:
    p = tmp_path / "guide.md"
    p.write_text(
        "---\ntitle: x\n---\n# Guide\n\nIntro text about the guide.\n\n## Install\n\nRun the installer.\n\n"
        "```\n# not a heading\n```\n\n- first\n- second\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n"
        "Setext Title\n------------\n\nBody under setext.\n", encoding="utf-8")
    return p


@pytest.fixture
def docx_styled(tmp_path: Path) -> Path:
    d = Document()
    d.add_heading("Annual Report", 0)
    d.add_heading("Finance", 1)
    d.add_paragraph(LOREM)
    d.add_heading("Revenue", 2)
    d.add_paragraph("Revenue reached 42 million euros in 2025.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Region", "Sales"
    t.cell(1, 0).text, t.cell(1, 1).text = "North", "10"
    d.save(tmp_path / "styled.docx")
    return tmp_path / "styled.docx"


@pytest.fixture
def docx_bold_only(tmp_path: Path) -> Path:
    """No heading styles at all: headings are bold Normal paragraphs (the ambiguous case)."""
    d = Document()
    for title, body in [("Safety rules", LOREM), ("Emergency exits", LOREM), ("Maintenance", LOREM)]:
        run = d.add_paragraph().add_run(title)
        run.bold = True
        run.font.size = Pt(11)
        p = d.add_paragraph(body)
        p.runs[0].font.size = Pt(11)
    d.save(tmp_path / "bold.docx")
    return tmp_path / "bold.docx"


@pytest.fixture
def pptx_file(tmp_path: Path) -> Path:
    prs = Presentation()
    for title, body, notes in [("Roadmap", "Ship v1 in March", "Mention budget"), ("Risks", "Vendor delay", "")]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    prs.save(tmp_path / "deck.pptx")
    return tmp_path / "deck.pptx"


@pytest.fixture
def xlsx_file(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Sales 2025"])
    ws.append(["Region", "Q1", "Q2"])
    ws.append(["North", 10, 12])
    ws.append(["South", 8, 9])
    ws.append([])
    ws.append(["Staff"])
    ws.append(["Name", "Role", "Age"])
    ws.append(["Ada", "Engineer", 36])
    ws.append(["Bob", "Analyst", 41])
    for c in ws[2] + ws[7]:
        c.font = Font(bold=True)
    wb.save(tmp_path / "book.xlsx")
    return tmp_path / "book.xlsx"


def make_pdf(path: Path, pages: int = 3, headings: bool = True) -> Path:
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    for page in range(1, pages + 1):
        c.setFont("Helvetica", 9)
        c.drawString(40, h - 25, "ACME Confidential Header")
        c.drawString(w / 2, 20, f"Page {page}")
        y = h - 90
        if headings:
            c.setFont("Helvetica-Bold", 20)
            c.drawString(50, y, f"Chapter {page}")
            y -= 34
            c.setFont("Helvetica-Bold", 14)
            c.drawString(50, y, f"Topic {page}.1")
            y -= 26
        c.setFont("Helvetica", 11)
        for line in range(6):
            c.drawString(50, y, f"Body line {line} of page {page} discussing the topic in detail here.")
            y -= 15
        c.showPage()
    c.save()
    return path


@pytest.fixture
def pdf_file(tmp_path: Path) -> Path:
    return make_pdf(tmp_path / "doc.pdf")


@pytest.fixture
def scanned_pdf(tmp_path: Path) -> Path:
    c = canvas.Canvas(str(tmp_path / "scan.pdf"), pagesize=A4)
    c.rect(50, 50, 300, 300, fill=1)
    c.showPage()
    c.save()
    return tmp_path / "scan.pdf"
