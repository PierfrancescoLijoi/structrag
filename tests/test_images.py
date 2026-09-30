import io
from dataclasses import replace
from pathlib import Path

import pytest

pytest.importorskip("rapidocr")
pytest.importorskip("pypdfium2")
from PIL import Image, ImageDraw, ImageFont   # noqa: E402

from structrag.app import build_services      # noqa: E402
from structrag.parsers import parse_file      # noqa: E402
from structrag.vision import ImageReader, get_reader   # noqa: E402


def _font(size: int):
    for name in ("arial.ttf", "DejaVuSans.ttf", "LiberationSans-Regular.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size)


def text_image(lines: list[str], size=(900, 300), font_size=40) -> Image.Image:
    img = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        d.text((30, 30 + i * (font_size + 20)), line, fill="black", font=_font(font_size))
    return img


@pytest.fixture()
def ocr_settings(settings):
    return replace(settings, ocr="auto")


def _png(path: Path, lines: list[str]) -> Path:
    text_image(lines).save(path)
    return path


def test_png_file_is_read_by_ocr_and_searchable(ocr_settings, tmp_path):
    f = _png(tmp_path / "targa_macchina.png", ["Maximum pressure 42 bar", "Service interval 900 hours"])
    sv = build_services(ocr_settings)
    r = sv.ingestor.ingest_file(f)
    assert r.status == "ingested"
    hit = sv.retriever.search("what is the service interval")[0]
    assert hit.kind == "image" and "900" in hit.text and "Text in figure" in hit.text
    assert (ocr_settings.data_dir / "images" / f"{hit.ref}.jpg").exists()          # thumbnail kept for citations
    ctx = sv.retriever.build_context([hit])[0]
    assert ctx.kind == "image" and ctx.ref == hit.ref


def test_scanned_pdf_is_ocrd_instead_of_flagged(ocr_settings, tmp_path):
    pages = [text_image([f"Chapter {n}", f"Boiler pressure is {n * 11} bar"]) for n in (1, 2)]
    pdf = tmp_path / "scan.pdf"
    pages[0].save(pdf, save_all=True, append_images=pages[1:])
    doc = parse_file(pdf, ocr_settings)
    text = " ".join(b.text for b in doc.blocks)
    assert not doc.needs_ocr and "22" in text and "11" in text
    assert {b.loc for b in doc.blocks} == {"p.1", "p.2"}
    assert any("OCR used for 2" in w for w in doc.warnings)


def test_scanned_pdf_without_ocr_is_still_flagged(settings, tmp_path):
    pdf = tmp_path / "scan.pdf"
    text_image(["Hidden text"]).save(pdf)
    assert parse_file(pdf, settings).needs_ocr            # ocr="off" in the shared fixture


def test_docx_inline_picture_is_indexed_once_even_if_repeated(ocr_settings, tmp_path):
    from docx import Document
    png = tmp_path / "fig.png"
    _png(png, ["Revenue 2025 was 42 million"])
    d = Document()
    d.add_heading("Report", 1)
    d.add_paragraph("Figure follows.")
    for _ in range(2):
        d.add_picture(str(png))
    tiny = tmp_path / "icon.png"
    Image.new("RGB", (20, 20), "red").save(tiny)
    d.add_picture(str(tiny))
    d.save(tmp_path / "r.docx")
    doc = parse_file(tmp_path / "r.docx", ocr_settings)
    figures = [b for b in doc.blocks if b.kind == "image"]
    assert len(figures) == 1 and "42 million" in figures[0].text        # duplicate and 20px icon skipped


def test_pptx_picture_is_read(ocr_settings, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    png = tmp_path / "chart.png"
    _png(png, ["Q3 growth 17 percent"])
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Results"
    slide.shapes.add_picture(str(png), Inches(1), Inches(2), width=Inches(6))
    prs.save(tmp_path / "d.pptx")
    doc = parse_file(tmp_path / "d.pptx", ocr_settings)
    assert any(b.kind == "image" and "17 percent" in b.text for b in doc.blocks)


def test_markdown_local_image_and_alt_fallback(ocr_settings, tmp_path):
    _png(tmp_path / "scheme.png", ["Valve V12 opens at 8 bar"])
    md = tmp_path / "n.md"
    md.write_text("# Notes\n\n![scheme](scheme.png)\n\n![remote logo](https://example.com/x.png)\n", encoding="utf-8")
    blocks = parse_file(md, ocr_settings).blocks
    assert any(b.kind == "image" and "8 bar" in b.text for b in blocks)
    assert any(b.text == "[Figure] remote logo" for b in blocks)


def test_ocr_result_is_cached_by_image_hash(ocr_settings, tmp_path, monkeypatch):
    calls = []
    reader = ImageReader(ocr_settings)
    original = reader._ocr_lines
    monkeypatch.setattr(reader, "_ocr_lines", lambda image: calls.append(1) or original(image))
    data = io.BytesIO()
    text_image(["Cached figure 77"]).save(data, "PNG")
    first, second = reader.read_bytes(data.getvalue()), reader.read_bytes(data.getvalue())
    assert first[1] == second[1] and "77" in first[1] and len(calls) == 1


def test_vision_model_caption_is_combined_with_ocr(ocr_settings, tmp_path, monkeypatch):
    s = replace(ocr_settings, vision_model="fake-vlm")
    reader = ImageReader(s)
    monkeypatch.setattr(reader, "_describe", lambda image: "A bar chart of quarterly revenue rising from 10 to 40.")
    data = io.BytesIO()
    text_image(["Revenue"]).save(data, "PNG")
    info, text = reader.read_bytes(data.getvalue())
    assert text.startswith("[Figure] A bar chart") and "Text in figure: Revenue" in text
    assert info.caption.startswith("A bar chart")


def test_vision_failure_falls_back_to_ocr_and_is_not_cached(ocr_settings, monkeypatch):
    s = replace(ocr_settings, vision_model="fake-vlm", vision_base_url="http://127.0.0.1:9/v1")   # nothing listens
    reader = ImageReader(s)
    data = io.BytesIO()
    text_image(["Only ocr here 5"]).save(data, "PNG")
    info, text = reader.read_bytes(data.getvalue())
    assert info.caption == "" and "5" in text and reader.cache.get(info.sha, "cap:fake-vlm") is None


def test_get_reader_is_shared_per_settings(ocr_settings):
    assert get_reader(ocr_settings) is get_reader(ocr_settings)
